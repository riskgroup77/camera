# -*- coding: utf-8 -*-
"""HEMIS'ga bog'lanmagan talabalarni HEMIS yozuvlari bilan birlashtirish.

Qoida va sabab — app/services/hemis_reconcile.py. Ikki bosqich:

  1. HISOBOT (bazaga tegmaydi). Excel fayl: "Birlashtiriladi", "Tekshirish
     kerak", "HEMIS'da topilmadi" varaqlari. Har qatorda "Qaror" ustuni:
       * "Birlashtiriladi" varag'ida — bo'sh qoldirilsa birlashtiriladi,
         "yo'q" deb yozilsa — tegilmaydi;
       * "Tekshirish kerak" varag'ida — faqat "ha" deb yozilganlari
         birlashtiriladi.
  2. QO'LLASH — o'sha Excel faylni o'qib, tasdiqlanganlarni birlashtiradi:
     HEMIS qatori qoladi (guruh, fakultet, dars jadvali), o'zi ro'yxatdan
     o'tgan qatorning yuzi, davomati, dars davomati, tashriflari va
     rozilik ma'lumoti unga ko'chadi, o'zi esa o'chiriladi
     (person_dedupe.merge_people). Har juftlik — alohida tranzaksiya va
     audit jurnalida yozuv. Oldin baza zaxirasini oling.

ISHGA TUSHIRISH (server, /opt/camera/camera-api):

    docker compose exec -T api python scripts/hemis_moslash.py hisobot --fayl /tmp/hemis_moslash.xlsx
    docker compose cp api:/tmp/hemis_moslash.xlsx ./hemis_moslash.xlsx
    # ... Excel'da ko'rib chiqiladi, keyin qaytariladi:
    docker compose cp ./hemis_moslash.xlsx api:/tmp/hemis_moslash.xlsx
    docker compose exec -T api python scripts/hemis_moslash.py qollash --fayl /tmp/hemis_moslash.xlsx --sinov
    docker compose exec -T api python scripts/hemis_moslash.py qollash --fayl /tmp/hemis_moslash.xlsx

Qayta ishga tushirish xavfsiz: birlashtirilgan qator ikkinchi marta topilmaydi.
"""

from __future__ import annotations

import argparse
import asyncio
import os
import sys
import uuid
from collections import Counter
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import text

from app.database import SessionLocal
from app.models import AuditLog
from app.services import hemis_reconcile as rec
from app.services import person_dedupe
from app.services.face_matching import announce_roster_change
from app.timezone import to_local

AUDIT_USER = "HEMIS moslashtirish"
SHEETS = {rec.MERGE: "Birlashtiriladi", rec.REVIEW: "Tekshirish kerak", rec.NOT_FOUND: "HEMISda topilmadi"}
DECISION = "Qaror"
KEEP_COL, DUP_COL = "hemis_qator_id", "ozi_qoshgan_qator_id"

_ROWS_SQL = """
    select s.id, s.full_name, s.type, s.pinfl, s.hemis_id, s.group_or_position, s.active,
           s.biometrics_status, s.biometric_embedding, s.biometric_photo_key,
           s.biometric_photo_left_key, s.biometric_photo_right_key, s.self_registered,
           s.consent_source, s.created_at, f.name as faculty,
           (select count(*) from attendance_records a where a.student_staff_id = s.id) as att
    from students_staff s left join faculties f on f.id = s.faculty_id
    where s.type = 'talaba' and (s.hemis_id is not null or s.active)
"""


async def load_rows(db) -> tuple[list[dict], list[dict]]:
    rows = [dict(r._mapping) for r in (await db.execute(text(_ROWS_SQL))).all()]
    unlinked = [r for r in rows if not r["hemis_id"] and r["active"]]
    hemis_rows = [r for r in rows if r["hemis_id"]]
    return unlinked, hemis_rows


def _face(row: dict | None) -> str:
    if not row or not row.get("biometric_embedding"):
        return "yo'q"
    angles = person_dedupe._all_angles(row)
    status = "tasdiqlangan" if row.get("biometrics_status") == "tasdiqlangan" else "kutilmoqda"
    return f"{status}{', 3 burchak' if angles else ''}"


def _date(value) -> str:
    return to_local(value).strftime("%d.%m.%Y") if isinstance(value, datetime) else ""


def _pinfl(row: dict | None) -> str:
    return f"…{row['pinfl'][-4:]}" if row and row.get("pinfl") else ""


def _person_cells(row: dict) -> list:
    return [
        row["full_name"], row["group_or_position"], _face(row), _pinfl(row), _date(row.get("created_at")),
        int(row.get("att") or 0), "o'zi (QR)" if row.get("self_registered") else (row.get("consent_source") or "admin"),
    ]


PERSON_HEAD = ["F.I.Sh. (bazada)", "Yozilgan guruhi", "Yuzi", "JSHSHIR", "Qo'shilgan", "Davomat kunlari", "Kim qo'shgan"]
HEMIS_HEAD = ["HEMIS F.I.Sh.", "HEMIS guruhi", "Fakultet", "HEMIS ID", "HEMIS'da yuzi", "HEMIS JSHSHIR", "Yuz o'xshashligi", "Sabab", "Nomzod"]


def _style(ws, widths: list[int]) -> None:
    head_fill = PatternFill("solid", fgColor="E8EEF9")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def build_workbook(proposals: list[rec.Proposal], unlinked_total: int) -> Workbook:
    wb = Workbook()
    summary = wb.active
    summary.title = "Xulosa"
    counts = Counter(p.verdict for p in proposals)
    summary.append(["HEMIS moslashtirish hisoboti", datetime.now().strftime("%d.%m.%Y %H:%M")])
    summary.append([])
    summary.append(["HEMIS'ga bog'lanmagan faol talabalar", unlinked_total])
    for verdict, label in rec.VERDICT_LABELS.items():
        summary.append([label, counts.get(verdict, 0)])
    summary.append([])
    summary.append(["Qanday ishlatiladi:"])
    summary.append(["• \"Birlashtiriladi\" — bo'sh qoldirilgan qatorlar birlashtiriladi; istamaganiga \"yo'q\" yozing."])
    summary.append(["• \"Tekshirish kerak\" — faqat \"ha\" yozilganlari birlashtiriladi."])
    summary.append(["• \"HEMISda topilmadi\" — avtomatik hech narsa qilinmaydi (qo'lda biriktiriladi yoki o'chiriladi)."])
    summary.append(["• Birlashtirishda HEMIS qatori qoladi; yuz, davomat va tashriflar unga ko'chadi."])
    summary["A1"].font = Font(bold=True, size=13)
    summary.column_dimensions["A"].width = 60
    summary.column_dimensions["B"].width = 18

    for verdict in (rec.MERGE, rec.REVIEW):
        ws = wb.create_sheet(SHEETS[verdict])
        ws.append(["№", DECISION, *PERSON_HEAD, *HEMIS_HEAD, KEEP_COL, DUP_COL])
        number = 0
        for p in proposals:
            if p.verdict != verdict:
                continue
            number += 1
            # Tekshirishda har nomzod — alohida qator: to'g'ri kelganiga "ha".
            options = [p.hemis] if p.hemis else []
            if verdict == rec.REVIEW:
                options += [r for r in p.alternatives if r is not p.hemis]
            for index, hemis in enumerate(options or [None], start=1):
                similarity = (p.face_similarity if hemis is p.hemis else rec.face_similarity(p.person, hemis)) if hemis else None
                ws.append([
                    number, "", *_person_cells(p.person),
                    hemis["full_name"] if hemis else "", hemis["group_or_position"] if hemis else "",
                    (hemis.get("faculty") or "") if hemis else "", hemis["hemis_id"] if hemis else "",
                    _face(hemis) if hemis else "", _pinfl(hemis),
                    similarity if similarity is not None else "", p.reason,
                    f"{index}/{len(options)}" if len(options) > 1 else "",
                    str(hemis["id"]) if hemis else "", str(p.person["id"]),
                ])
        _style(ws, [5, 9, 32, 22, 18, 10, 11, 9, 11, 32, 30, 24, 13, 16, 10, 9, 48, 9, 38, 38])

    ws = wb.create_sheet(SHEETS[rec.NOT_FOUND])
    ws.append(["№", *PERSON_HEAD, "Izoh", DUP_COL])
    number = 0
    for p in proposals:
        if p.verdict == rec.NOT_FOUND:
            number += 1
            ws.append([number, *_person_cells(p.person), p.reason, str(p.person["id"])])
    _style(ws, [5, 32, 22, 18, 10, 11, 9, 11, 32, 38])
    return wb


async def cmd_hisobot(path: str) -> None:
    async with SessionLocal() as db:
        unlinked, hemis_rows = await load_rows(db)
    proposals = rec.reconcile(unlinked, hemis_rows)
    build_workbook(proposals, len(unlinked)).save(path)
    counts = Counter(p.verdict for p in proposals)
    print(f"HEMIS'ga bog'lanmagan faol talabalar: {len(unlinked)} (HEMIS qatorlari: {len(hemis_rows)})")
    for verdict, label in rec.VERDICT_LABELS.items():
        print(f"  {label:<22}: {counts.get(verdict, 0)}")
    print(f"Hisobot: {path}")


def read_decisions(path: str) -> list[tuple[uuid.UUID, uuid.UUID, str]]:
    """(hemis qatori, o'zi qo'shgan qator, varaq) — birlashtiriladiganlar."""
    wb = load_workbook(path, read_only=True)
    chosen: list[tuple[uuid.UUID, uuid.UUID, str]] = []
    for verdict in (rec.MERGE, rec.REVIEW):
        title = SHEETS[verdict]
        if title not in wb.sheetnames:
            continue
        rows = wb[title].iter_rows(values_only=True)
        head = list(next(rows, []) or [])
        try:
            i_dec, i_keep, i_dup = head.index(DECISION), head.index(KEEP_COL), head.index(DUP_COL)
        except ValueError:
            raise SystemExit(f"'{title}' varag'ida ustunlar o'zgartirilgan — yangi hisobot oling")
        for row in rows:
            if not row or not row[i_keep] or not row[i_dup]:
                continue
            decision = str(row[i_dec] or "").strip().lower().replace("‘", "'").replace("’", "'")
            take = decision in ("ha", "xa", "yes", "+") if verdict == rec.REVIEW else decision not in ("yo'q", "yoq", "no", "-")
            if take:
                chosen.append((uuid.UUID(str(row[i_keep])), uuid.UUID(str(row[i_dup])), title))
    return chosen


async def cmd_qollash(path: str, dry_run: bool) -> None:
    pairs = read_decisions(path)
    keepers = Counter(keep for keep, _, _ in pairs)
    dups = Counter(dup for _, dup, _ in pairs)
    # Bitta HEMIS qatoriga bir nechta yozuv — faqat ular o'zaro bitta odam
    # bo'lsa (ikki marta ro'yxatdan o'tgan); aks holda hech biri.
    shared = {keep for keep, n in keepers.items() if n > 1}
    if shared:
        async with SessionLocal() as db:
            rows = {
                r["id"]: r for r in (dict(x._mapping) for x in (await db.execute(
                    text("select id, full_name, pinfl, biometric_embedding from students_staff where id = any(:ids)"),
                    {"ids": [dup for keep, dup, _ in pairs if keep in shared]},
                )).all())
            }
        for keep in list(shared):
            people = [rows[dup] for k, dup, _ in pairs if k == keep and dup in rows]
            if rec.mutually_same(people):
                shared.discard(keep)
    print(f"Faylda birlashtiriladigan juftliklar: {len(pairs)}{'  (SINOV — bazaga yozilmaydi)' if dry_run else ''}")
    merged = 0
    totals: Counter = Counter()
    errors: list[str] = []
    for keep_id, dup_id, sheet in pairs:
        if keep_id in shared:
            errors.append(f"{keep_id}: bitta HEMIS qatoriga bir-biridan farqli odamlar tanlangan — o'tkazildi")
            continue
        if dups[dup_id] > 1:
            errors.append(f"{dup_id}: bitta odamga bir nechta HEMIS nomzodi tanlangan — o'tkazildi")
            continue
        async with SessionLocal() as db:
            check = (
                await db.execute(
                    text("select id, hemis_id, full_name from students_staff where id in (:k, :d)"),
                    {"k": keep_id, "d": dup_id},
                )
            ).all()
            found = {row.id: row for row in check}
            keep, dup = found.get(keep_id), found.get(dup_id)
            if keep is None or dup is None:
                errors.append(f"{dup_id}: yozuv topilmadi (allaqachon birlashtirilgan?)")
                continue
            if not keep.hemis_id or dup.hemis_id:
                errors.append(f"{dup.full_name}: HEMIS bog'lanishi o'zgargan — yangi hisobot oling")
                continue
            try:
                moved = await person_dedupe.merge_people(db, keep_id, dup_id)
            except person_dedupe.MergeError as exc:
                await db.rollback()
                errors.append(f"{dup.full_name}: {exc}")
                continue
            detail = ", ".join(f"{k}: {v}" for k, v in moved.items()) or "bo'sh"
            db.add(AuditLog(
                user_id=None, user_name=AUDIT_USER, module="Talabalar", status="muvaffaqiyatli", ip="internal",
                action=f"HEMIS bilan birlashtirildi: {keep.full_name} ({keep.hemis_id}) ← {dup.full_name} ({detail}) [{sheet}]",
            ))
            if dry_run:
                await db.rollback()
            else:
                await db.commit()
            merged += 1
            totals.update(moved)
            print(f"  ✓ {dup.full_name} → {keep.full_name} ({keep.hemis_id}): {detail}")
    print(f"\nBirlashtirildi: {merged}{' (sinov)' if dry_run else ''}; ko'chirilgan: {dict(totals)}")
    for error in errors:
        print(f"  ! {error}")
    if merged and not dry_run:
        await announce_roster_change()
        print("Tanish ro'yxati yangilanishi e'lon qilindi.")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest="cmd", required=True)
    report = sub.add_parser("hisobot", help="moslashtirish hisoboti (bazaga tegmaydi)")
    report.add_argument("--fayl", default="/tmp/hemis_moslash.xlsx")
    apply = sub.add_parser("qollash", help="Excel'dagi tasdiqlanganlarni birlashtirish")
    apply.add_argument("--fayl", required=True)
    apply.add_argument("--sinov", action="store_true", help="hamma tekshiruvdan o'tadi, lekin bazaga yozilmaydi")
    args = parser.parse_args()
    if args.cmd == "hisobot":
        asyncio.run(cmd_hisobot(args.fayl))
    else:
        asyncio.run(cmd_qollash(args.fayl, args.sinov))


if __name__ == "__main__":
    main()
