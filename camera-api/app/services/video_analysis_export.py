"""Kunlik video tahlil natijalari — Excel (.xlsx).

Uch varaq: "Natijalar" (har odam — barcha kriteriyalar bitta qatorda),
"Xulosa" (kriteriyalar bo'yicha sonlar) va "Holatlar" (har aniqlangan holat,
vaqti, kamerasi va video dalili bor-yo'qligi). API (`/api/video-tahlil/natijalar/
export.xlsx`) va buyruq qatori (scripts/video_tahlil.py) bir xil faylni beradi.
"""

from __future__ import annotations

from datetime import date as date_type, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from app.services.staff_export import _header, _landscape, _stamp, _title, _to_bytes, _widths
from app.timezone import local_now, to_local

ATTENDANCE = {"keldi": "Keldi", "kech_keldi": "Kech keldi", "kelmadi": "Kelmadi"}
EARLY = {"erta_ketdi": "Erta ketdi", "vaqtida": "Vaqtida", "aniqlanmadi": "Aniqlanmadi", "tegishli_emas": ""}
COAT = {"kiygan": "Kiygan", "kiymagan": "Kiymagan", "aniqlanmadi": "Aniqlanmadi", "talab_yoq": "Talab yo'q"}

_RED = PatternFill("solid", fgColor="FDE2E1")
_AMBER = PatternFill("solid", fgColor="FFF1D6")
_THIN = Side(style="thin", color="D5DAE1")
_CELL_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

HEADERS = [
    "F.I.Sh.", "Turi", "Guruh / lavozim", "Davomat", "Kelgan", "Ketgan", "Kechikish, daq", "Ishdan erta ketish",
    "Darslar (kirgan/o'lchangan)", "Darsga kech", "Darsdan erta", "Diqqat (0-100)", "Oq xalat",
    "Oq xalat (oq/kuzatuv)", "Chekish", "O'qituvchi: darslar", "Vaqtida", "Kech", "Kelmadi", "Faollik (0-100)",
]
WIDTHS = [34, 9, 24, 12, 9, 9, 11, 14, 14, 10, 10, 11, 12, 12, 9, 12, 9, 7, 9, 11]


def _clock(moment: datetime | None) -> str:
    return to_local(moment).strftime("%H:%M") if moment else ""


def _lessons(row) -> str:
    if not row.lessons_total:
        return ""
    measured = row.lessons_total - (row.lessons_unmeasured or 0)
    return f"{row.lessons_attended or 0}/{measured}" if measured > 0 else "o'lchanmadi"


def row_values(row, person) -> list:
    return [
        person.full_name,
        "Talaba" if person.type == "talaba" else "Xodim",
        person.group_or_position,
        ATTENDANCE.get(row.attendance_status or "", ""),
        _clock(row.arrived_at),
        _clock(row.left_at),
        row.late_minutes or "",
        EARLY.get(row.early_leave or "", ""),
        _lessons(row),
        row.lessons_late if row.lessons_total else "",
        row.lessons_left_early if row.lessons_total else "",
        row.attention_score if row.attention_score is not None else "",
        COAT.get(row.coat_status or "", ""),
        f"{row.coat_white_samples}/{row.coat_samples}" if row.coat_samples else "",
        row.smoking_events or "",
        row.teacher_lessons or "",
        row.teacher_on_time if row.teacher_lessons else "",
        row.teacher_late if row.teacher_lessons else "",
        row.teacher_absent if row.teacher_lessons else "",
        row.teacher_activity if row.teacher_activity is not None else "",
    ]


def _flag(row) -> PatternFill | None:
    if row.attendance_status == "kelmadi" or row.coat_status == "kiymagan" or row.smoking_events:
        return _RED
    if (row.late_minutes or 0) > 0 or row.early_leave == "erta_ketdi" or (row.lessons_late or 0) > 0:
        return _AMBER
    return None


def build_criteria_workbook(
    day: date_type, rows: list[tuple], *, title_suffix: str = "", camera_names: dict[str, str] | None = None
) -> bytes:
    """`rows` — (DailyPersonCriteria, StudentStaff) juftliklari. Uchinchi
    varaq "Holatlar" — har aniqlangan holat (dalil bilan) alohida qatorda."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Natijalar"
    _title(
        ws,
        f"Kunlik video tahlil — {day.strftime('%d.%m.%Y')}{title_suffix}",
        f"NVR yozuvlaridan hisoblangan kriteriyalar. Bo'sh katak — o'lchanmadi (kamera ko'rmagan yoki talab "
        f"qilinmaydi). Tuzildi: {_stamp(local_now())}",
        len(HEADERS),
    )
    _header(ws, 4, HEADERS)
    for index, (row, person) in enumerate(sorted(rows, key=lambda item: item[1].full_name), start=5):
        fill = _flag(row)
        for col, value in enumerate(row_values(row, person), start=1):
            cell = ws.cell(row=index, column=col, value=value)
            cell.border = _CELL_BORDER
            cell.alignment = Alignment(vertical="center", horizontal="left" if col <= 3 else "center")
            if fill is not None:
                cell.fill = fill
    _widths(ws, WIDTHS)
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"A4:{ws.cell(row=4, column=len(HEADERS)).column_letter}{max(5, len(rows) + 4)}"
    _landscape(ws, 4)

    summary = wb.create_sheet("Xulosa")
    data = [r for r, _p in rows]
    stats = [
        ("Odamlar (natijasi bor)", len(data)),
        ("Keldi", sum(1 for r in data if r.attendance_status in ("keldi", "kech_keldi"))),
        ("Kech keldi", sum(1 for r in data if r.attendance_status == "kech_keldi")),
        ("Kelmadi", sum(1 for r in data if r.attendance_status == "kelmadi")),
        ("Ishdan erta ketdi", sum(1 for r in data if r.early_leave == "erta_ketdi")),
        ("Darsga kechikish (talaba-dars)", sum(r.lessons_late or 0 for r in data)),
        ("Darsdan erta chiqish (talaba-dars)", sum(r.lessons_left_early or 0 for r in data)),
        ("O'rtacha diqqat", _avg([r.attention_score for r in data])),
        ("Oq xalat kiygan", sum(1 for r in data if r.coat_status == "kiygan")),
        ("Oq xalat kiymagan", sum(1 for r in data if r.coat_status == "kiymagan")),
        ("Chekish holatlari", sum(r.smoking_events or 0 for r in data)),
        ("O'qituvchi darslari: vaqtida", sum(r.teacher_on_time or 0 for r in data)),
        ("O'qituvchi darslari: kech", sum(r.teacher_late or 0 for r in data)),
        ("O'qituvchi darslari: kelmadi", sum(r.teacher_absent or 0 for r in data)),
        ("O'qituvchi faolligi (o'rtacha)", _avg([r.teacher_activity for r in data])),
    ]
    summary["A1"] = f"Xulosa — {day.strftime('%d.%m.%Y')}"
    summary["A1"].font = Font(bold=True, size=14)
    for i, (label, value) in enumerate(stats, start=3):
        summary.cell(row=i, column=1, value=label)
        summary.cell(row=i, column=2, value=value if value is not None else "—")
    summary.column_dimensions["A"].width = 38
    summary.column_dimensions["B"].width = 12

    _findings_sheet(wb, rows, camera_names or {})
    return _to_bytes(wb)


CRITERIA = {
    6: "Xodim davomati", 7: "Talaba davomati", 8: "Darsga kechikish", 9: "Darsdan/ishdan erta ketish",
    10: "Oq xalat", 15: "Chekish", 19: "Darsga diqqat", 21: "O'qituvchi faolligi", 22: "O'qituvchining darsga kelishi",
}
FINDING_HEADERS = ["F.I.Sh.", "Turi", "Guruh / lavozim", "Kriteriya", "Holat", "Vaqt", "Kamera", "Video dalil (2 daq)"]


def _findings_sheet(wb: Workbook, rows: list[tuple], camera_names: dict[str, str]) -> None:
    """Har aniqlangan holat alohida qatorda — video dalil saytda odam kartasida."""
    ws = wb.create_sheet("Holatlar")
    _title(ws, "Aniqlangan holatlar", "Har holat uchun rasm va 2 daqiqalik video dalil saytda: Kunlik tahlil → Natijalar → odam", len(FINDING_HEADERS))
    _header(ws, 4, FINDING_HEADERS)
    line = 5
    for row, person in sorted(rows, key=lambda item: item[1].full_name):
        for item in (row.details or {}).get("dalillar") or []:
            try:
                moment = datetime.fromisoformat(item["vaqt"])
            except (KeyError, TypeError, ValueError):
                continue
            code = int(item.get("kod") or 0)
            values = [
                person.full_name,
                "Talaba" if person.type == "talaba" else "Xodim",
                person.group_or_position,
                f"{code}. {CRITERIA.get(code, '')}",
                item.get("sabab") or "",
                to_local(moment).strftime("%H:%M:%S"),
                camera_names.get(str(item.get("kamera_id")), ""),
                "bor" if item.get("klip") else (f"yo'q: {item['klip_xato']}" if item.get("klip_xato") else "kutilmoqda"),
            ]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=line, column=col, value=value)
                cell.border = _CELL_BORDER
                cell.alignment = Alignment(vertical="center", wrap_text=col == 5)
            line += 1
    _widths(ws, [30, 9, 22, 26, 60, 10, 24, 18])
    ws.freeze_panes = "B5"
    _landscape(ws, 4)


def build_findings_html(day: date_type, rows: list[tuple], camera_names: dict[str, str], url_for) -> str:
    """Holatlar va video dalillar — bitta HTML sahifa (sinov natijalarini
    saytsiz ko'rish uchun). `url_for(key)` — klip havolasi (presigned)."""
    from html import escape

    people = []
    totals: dict[int, int] = {}
    for row, person in sorted(rows, key=lambda item: item[1].full_name):
        items = (row.details or {}).get("dalillar") or []
        if not items:
            continue
        cards = []
        for item in items:
            try:
                moment = to_local(datetime.fromisoformat(item["vaqt"])).strftime("%H:%M:%S")
            except (KeyError, TypeError, ValueError):
                continue
            code = int(item.get("kod") or 0)
            totals[code] = totals.get(code, 0) + 1
            if item.get("klip"):
                media = f'<video controls preload="none" src="{escape(url_for(item["klip"]))}"></video>'
            else:
                media = f'<p class="muted">Video dalil yo\'q{": " + escape(item["klip_xato"]) if item.get("klip_xato") else ""}</p>'
            cards.append(
                f'<li class="card"><div class="head"><span class="tag">{code}. {escape(CRITERIA.get(code, ""))}</span>'
                f'<span class="muted">{moment} · {escape(camera_names.get(str(item.get("kamera_id")), ""))}</span></div>'
                f"<p>{escape(item.get('sabab') or '')}</p>{media}</li>"
            )
        people.append(
            f'<section><h2>{escape(person.full_name)} <small>{escape(person.group_or_position or "")} · '
            f'{"talaba" if person.type == "talaba" else "xodim"}</small></h2><ul>{"".join(cards)}</ul></section>'
        )
    summary = "".join(
        f"<tr><td>{code}. {escape(CRITERIA.get(code, ''))}</td><td>{count}</td></tr>" for code, count in sorted(totals.items())
    )
    return f"""<!doctype html><html lang="uz"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1"><title>Holatlar — {day:%d.%m.%Y}</title>
<style>
:root{{--bg:#f6f7f9;--fg:#18202b;--muted:#5d6878;--card:#fff;--line:#dfe3e8;--tag:#b42318;--tagbg:#fde2e1}}
@media (prefers-color-scheme:dark){{:root{{--bg:#11151b;--fg:#e7ebf0;--muted:#9aa4b2;--card:#1a2029;--line:#2a313c;--tag:#ffb4ab;--tagbg:#4a1d1a}}}}
body{{margin:0;padding:16px;background:var(--bg);color:var(--fg);font:14px/1.5 system-ui,sans-serif}}
main{{max-width:1100px;margin:0 auto}} h1{{font-size:22px}} h2{{font-size:16px;margin:24px 0 8px}}
small,.muted{{color:var(--muted);font-weight:normal}} table{{border-collapse:collapse;margin:8px 0 16px}}
td{{border:1px solid var(--line);padding:4px 10px}} ul{{list-style:none;padding:0;margin:0;display:grid;gap:12px;grid-template-columns:repeat(auto-fill,minmax(300px,1fr))}}
.card{{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:12px;min-width:0}}
.head{{display:flex;flex-wrap:wrap;gap:8px;align-items:center}} .tag{{background:var(--tagbg);color:var(--tag);border-radius:6px;padding:1px 8px;font-weight:600}}
video{{width:100%;aspect-ratio:16/9;background:#000;border-radius:6px}} p{{margin:8px 0}}
</style></head><body><main>
<h1>Aniqlangan holatlar — {day:%d.%m.%Y}</h1>
<p class="muted">Har holat uchun 2 daqiqalik video dalil (holatdan 1 daqiqa oldin va keyin). Havolalar 7 kun amal qiladi.</p>
<table>{summary}</table>
{"".join(people) or "<p>Holat aniqlanmadi.</p>"}
</main></body></html>"""


def _avg(values: list) -> int | None:
    present = [v for v in values if v is not None]
    return round(sum(present) / len(present)) if present else None
