"""Avtomatik hisobot yuborish (report_schedules jadvali).

Rahbar har dushanba ertalab o'tgan hafta, har oyning 1-sanasida o'tgan oy
hisobotini Telegram'da Excel fayl ko'rinishida oladi — sahifaga kirib
yuklab olishi shart emas. Fayl mavjud quruvchilar bilan tuziladi (Hisobotlar
sahifasidagi "Excel" tugmasi bilan bir xil fayl), izohda esa KPI ning
asosiy raqamlari: faylni ochmasdan ham holat ko'rinsin.

Qachon yuboriladi: davr chegarasi (haftalik — dushanba 08:00, oylik —
1-sana 08:00, Toshkent vaqti) o'tgan va shu chegaradan keyin hali
yuborilmagan bo'lsa. Server dushanba kuni o'chiq bo'lsa ham, keyin
yoqilganda o'sha hafta ichida yetkaziladi. Chegaradan KEYIN yaratilgan
jadval o'tgan davrni yubormaydi — yangi qo'shilgan jadval darhol "eski"
hisobot jo'natib yubormasin; keyingi chegarani kutadi (darhol kerak
bo'lsa — "Sinov" tugmasi).
"""

from __future__ import annotations

import html
import logging
from dataclasses import dataclass
from datetime import date as date_type, datetime, time, timedelta, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from app.config import settings
from app.models import NotificationLog, ReportSchedule
from app.services import hisobot, kpi as kpi_svc, tabel as tabel_svc
from app.services.attendance_policy import load_policy
from app.services.notifications import telegram
from app.timezone import INSTITUTE_TZ

logger = logging.getLogger("app.report_schedule")

KINDS = {"kunlik": "Kunlik", "haftalik": "Haftalik", "oylik": "Oylik"}
REPORTS = {
    "kpi": "KPI",
    "davomat_xodim": "Xodimlar davomati",
    "davomat_talaba": "Talabalar davomati",
    "tabel_xodim": "Xodimlar tabeli",
    "tabel_talaba": "Talabalar tabeli",
    "jadval_davomat": "Jadval bo'yicha davomat",
}
SEND_HOUR = 8
# Kunlik hisobot — o'sha kunning o'zi, "kelmadi" belgilangandan KEYIN
# (attendance_absence_mark_after, standart 20:00). Ilgari 19:00 da ketardi —
# kelmaganlar hali yozilmagan, hisobotda doim "0 kelmadi" chiqardi.
DAILY_SEND_HOUR = 21
LOG_KIND = "report"


# ─────────────────────────────────────────── davrlar

def _at(day: date_type) -> datetime:
    return datetime.combine(day, time(SEND_HOUR), tzinfo=INSTITUTE_TZ)


def _month_start(day: date_type) -> date_type:
    return day.replace(day=1)


def _prev_month_start(day: date_type) -> date_type:
    return _month_start(_month_start(day) - timedelta(days=1))


def last_boundary(kind: str, now: datetime) -> datetime:
    """`now` dan oldingi (yoki unga teng) eng so'nggi yuborish chegarasi."""
    local = now.astimezone(INSTITUTE_TZ)
    if kind == "kunlik":
        boundary = datetime.combine(local.date(), time(DAILY_SEND_HOUR), tzinfo=INSTITUTE_TZ)
        return boundary if boundary <= local else boundary - timedelta(days=1)
    if kind == "haftalik":
        boundary = _at(local.date() - timedelta(days=local.weekday()))
        return boundary if boundary <= local else boundary - timedelta(days=7)
    boundary = _at(_month_start(local.date()))
    return boundary if boundary <= local else _at(_prev_month_start(local.date()))


def period_for(kind: str, boundary: datetime) -> tuple[date_type, date_type]:
    """Chegaradan oldingi to'liq davr: o'tgan hafta (Du–Ya) yoki o'tgan oy."""
    day = boundary.astimezone(INSTITUTE_TZ).date()
    if kind == "kunlik":
        return day, day
    if kind == "haftalik":
        return day - timedelta(days=7), day - timedelta(days=1)
    return _prev_month_start(day), day - timedelta(days=1)


def due_period(schedule: ReportSchedule, now: datetime) -> tuple[date_type, date_type] | None:
    if not schedule.enabled or not schedule.telegram_chat_ids:
        return None
    boundary = last_boundary(schedule.kind, now)
    if schedule.last_sent_at is not None and schedule.last_sent_at >= boundary:
        return None
    if schedule.last_sent_at is None and schedule.created_at is not None and schedule.created_at > boundary:
        return None
    return period_for(schedule.kind, boundary)


def period_label(start: date_type, end: date_type) -> str:
    return f"{start:%d.%m.%Y} — {end:%d.%m.%Y}"


# ─────────────────────────────────────────── fayl

def build_kpi_workbook(data: dict) -> bytes:
    """KPI paneli bitta varaqda: blok nomi, ko'rsatkich, qiymat."""
    wb = Workbook()
    ws = wb.active
    ws.title = "KPI"
    bold = Font(bold=True)
    period = data["period"]
    ws.append([settings.org_name])
    ws.append(["Rahbariyat KPI"])
    ws["A2"].font = Font(bold=True, size=14)
    ws.append(["Davr", f"{period['from']} — {period['to']}"])
    ws.append([])

    def section(title: str, rows: list[tuple[str, object]]) -> None:
        ws.append([title])
        ws.cell(ws.max_row, 1).font = bold
        for label, value in rows:
            ws.append([label, "—" if value is None else value])
        ws.append([])

    att, rec, sec, inf = data["attendance"], data["recognition"], data["security"], data["infrastructure"]
    section("Davomat, %", [
        ("Xodimlar", att["staff"]["rate"]), ("Xodimlar (oldingi davr)", att["staff"]["prevRate"]),
        ("Talabalar", att["students"]["rate"]), ("Talabalar (oldingi davr)", att["students"]["prevRate"]),
        ("Kechikish — xodimlar", att["staff"]["latePct"]), ("Kechikish — talabalar", att["students"]["latePct"]),
    ])
    section("Yuzni tanish", [
        ("Yuz bazasi — xodimlar, %", rec["staffCoverage"]), ("Yuz bazasi — talabalar, %", rec["studentsCoverage"]),
        ("Kuniga tanilgan, %", rec["recognisedDailyPct"]), ("Notanish yuzlar (kutmoqda)", rec["unknownPending"]),
    ])
    section("Xavfsizlik", [
        ("Hodisalar", sec["events"]), ("Ko'rib chiqilgan", sec["reviewed"]), ("Yolg'on signal, %", sec["falsePct"]),
        ("Javob vaqti, daq", sec["reviewMinutes"]), ("Hal qilish vaqti, daq", sec["resolveMinutes"]),
    ])
    section("Kameralar", [
        ("Jami", inf["camerasTotal"]), ("Faol", inf["camerasActive"]), ("Ishlayapti", inf["camerasOnline"]),
        ("Ishlayapti, %", inf["onlinePct"]),
    ])
    if sec["modules"]:
        ws.append(["Modullar bo'yicha"])
        ws.cell(ws.max_row, 1).font = bold
        ws.append(["Modul", "Hodisalar", "Ko'rilgan", "Yolg'on, %", "Javob, daq", "Hal qilish, daq"])
        for cell in ws[ws.max_row]:
            cell.font = bold
        for m in sec["modules"]:
            ws.append([m["name"], m["events"], m["reviewed"], m["falsePct"], m["reviewMinutes"], m["resolveMinutes"]])
    ws.column_dimensions["A"].width = 36
    for col in "BCDEF":
        ws.column_dimensions[col].width = 14
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


@dataclass
class Attachment:
    filename: str
    content: bytes
    caption: str


async def build_attachment(db: AsyncSession, schedule: ReportSchedule, start: date_type, end: date_type) -> Attachment:
    # Fayl quruvchilari router modulida (Excel tugmasi bilan bir xil fayl).
    from app.routers.hisobot import build_tabel_workbook, build_workbook

    kpi = await kpi_svc.build(db, start, end)
    report = schedule.report
    stamp = f"{start}_{end}"
    if report == "jadval_davomat":
        from app.services.schedule_presence import board_workbook, day_board

        board = await day_board(db, end)
        content, filename = board_workbook(end, board), f"jadval-davomat-{end.isoformat()}.xlsx"
        not_seen = sum(1 for row in board if row.teacher_status == "kelmagan")
        expected = sum(row.students_expected for row in board)
        arrived = sum(row.students_arrived for row in board)
        lines = [
            f"<b>{html.escape(schedule.name)}</b>",
            f"{REPORTS[report]} · {end:%d.%m.%Y}",
            "",
            f"Darslar: {len(board)}",
            f"O'qituvchini kamera ko'rmadi: {not_seen}",
            f"Talabalar keldi: {arrived}/{expected}" if expected else "Talabalar: —",
        ]
        return Attachment(filename=filename, content=content, caption="\n".join(lines))
    if report == "kpi":
        content, filename = build_kpi_workbook(kpi), f"kpi-{stamp}.xlsx"
    elif report.startswith("davomat_"):
        kind = report.removeprefix("davomat_")
        data = await hisobot.report(db, kind, start, end, hisobot.Filters(), "davomat", limit=None)
        content = build_workbook(data)
        filename = f"davomat-{'talabalar' if kind == 'talaba' else 'xodimlar'}-{stamp}.xlsx"
    else:
        kind = report.removeprefix("tabel_")
        # Tabel oylik: haftalik jadvalda hafta tugagan oy olinadi.
        data = await tabel_svc.build(db, kind, f"{end.year:04d}-{end.month:02d}", hisobot.Filters())
        content = build_tabel_workbook(data)
        filename = f"tabel-{'talabalar' if kind == 'talaba' else 'xodimlar'}-{data['month']}.xlsx"
    lines = [f"<b>{html.escape(schedule.name)}</b>", f"{REPORTS[report]} · {period_label(start, end)}", ""]
    lines += [html.escape(line) for line in kpi_svc.summary_lines(kpi)]
    return Attachment(filename=filename, content=content, caption="\n".join(lines))


@dataclass
class SendOutcome:
    sent: int
    failed: int
    errors: list[str]
    all_blocked: bool
    # Javob kelmagan (timeout) — fayl yetib borgan bo'lishi mumkin.
    uncertain: int = 0


async def send(db: AsyncSession, schedule: ReportSchedule, start: date_type, end: date_type) -> SendOutcome:
    """Faylni har bir chatga yuboradi va natijani notification_log'ga yozadi
    ("Bildirishnomalar" jurnalida ko'rinsin). Commit chaqiruvchida."""
    attachment = await build_attachment(db, schedule, start, end)
    sent = failed = blocked = uncertain = 0
    errors: list[str] = []
    text = telegram.strip_html(attachment.caption)
    for chat_id in schedule.telegram_chat_ids or []:
        result = await telegram.send_document(chat_id, attachment.content, attachment.filename, attachment.caption)
        if result.ok:
            sent += 1
        else:
            failed += 1
            blocked += result.blocked
            uncertain += result.no_fallback
            errors.append(f"{chat_id}: {result.error}")
        db.add(NotificationLog(
            channel="telegram", recipient=str(chat_id), kind=LOG_KIND,
            status="yuborildi" if result.ok else "xato", text=text, error=result.error, ref_id=str(schedule.id),
        ))
    return SendOutcome(
        sent=sent, failed=failed, errors=errors, all_blocked=failed > 0 and blocked == failed, uncertain=uncertain
    )


#: Bir davr uchun xato bilan tugagan urinishlar chegarasi.
MAX_ATTEMPTS = 3


async def _failed_attempts(db: AsyncSession, schedule: ReportSchedule, now: datetime) -> int:
    """Shu davr chegarasidan beri nechta urinish xato bilan tugagan (jurnal bo'yicha)."""
    from sqlalchemy import func

    boundary = last_boundary(schedule.kind, now)
    return int(
        await db.scalar(
            select(func.count(func.distinct(NotificationLog.created_at)))
            .where(NotificationLog.kind == LOG_KIND, NotificationLog.ref_id == str(schedule.id))
            .where(NotificationLog.status == "xato", NotificationLog.created_at >= boundary)
        )
        or 0
    )


async def run_due(session_factory: async_sessionmaker[AsyncSession], now: datetime | None = None) -> int:
    """Muddati kelgan jadvallarni yuboradi; yuborilgan jadvallar sonini qaytaradi."""
    if not telegram.is_configured():
        return 0
    now = now or datetime.now(timezone.utc)
    done = 0
    async with session_factory() as db:
        schedules = (await db.execute(select(ReportSchedule).where(ReportSchedule.enabled.is_(True)))).scalars().all()
        due = [(s.id, period) for s in schedules if (period := due_period(s, now)) is not None]
    for schedule_id, period in due:
        # Har jadval o'z sessiyasida: bittasi (masalan, bo'sh ma'lumot)
        # yiqilsa qolganlari to'xtamasin.
        async with session_factory() as db:
            schedule = await db.get(ReportSchedule, schedule_id)
            if schedule is None:
                continue
            if schedule.kind == "kunlik" and not (await load_policy(db)).is_work_day(period[0]):
                # Dam olish kuni / bayram — kunlik davomat hisoboti bo'sh, yuborilmaydi.
                schedule.last_sent_at = now
                await db.commit()
                continue
            try:
                outcome = await send(db, schedule, *period)
            except Exception:
                logger.exception("report schedule failed", extra={"schedule": str(schedule_id)})
                await db.rollback()
                continue
            # Vaqtinchalik xato (tarmoq) bo'lsa keyingi aylanishda qayta
            # urinamiz; bloklangan chatga qayta urinishdan foyda yo'q.
            # Timeout — fayl yetib borgan bo'lishi mumkin: qayta yuborilsa rahbariyat
            # bir hisobotni har 5 daqiqada olardi. Doimiy xato (masalan chat
            # supergroup'ga aylangan) — shu davr uchun ko'pi bilan MAX_ATTEMPTS marta.
            attempts = await _failed_attempts(db, schedule, now)
            if outcome.sent or outcome.all_blocked or outcome.uncertain or attempts >= MAX_ATTEMPTS:
                schedule.last_sent_at = now
                done += 1
            await db.commit()
    return done
