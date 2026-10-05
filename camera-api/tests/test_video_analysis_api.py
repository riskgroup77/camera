"""Kunlik video tahlil API: holat, ishga tushirish/bekor qilish, natijalar,
NVR sozlamasi va huquqlar."""

import shutil
import subprocess
import uuid
from datetime import date, datetime, time, timedelta, timezone

import pytest
from sqlalchemy import select

from app.models import (
    AuditLog,
    Camera,
    DailyPersonCriteria,
    NvrDevice,
    StudentStaff,
    VideoAnalysisRun,
)
from app.timezone import business_today
from tests.conftest import auth_headers

pytestmark = pytest.mark.daily_mode


@pytest.fixture
async def admin(client, seeded):
    return await auth_headers(client, "admin", "admin123")


@pytest.fixture
async def operator(client, seeded):
    return await auth_headers(client, "operator", "operator123")


async def test_status_in_daily_mode(client, admin, db_session):
    db_session.add(NvrDevice(name="N", kind="fayl", base_path="/x"))
    await db_session.commit()
    resp = await client.get("/api/video-tahlil/holat", headers=admin)
    assert resp.status_code == 200
    body = resp.json()
    assert body["mode"] == "kunlik" and body["startTime"] == "20:00"
    assert body["nextRunAt"] and body["nvrCount"] == 1
    assert body["current"] is None


async def test_live_overlay_is_off_in_daily_mode(client, admin, db_session):
    camera = Camera(name="K", ip="10.1.1.1", zone="A", resolution="HD", status="faol")
    db_session.add(camera)
    await db_session.commit()
    resp = await client.get(f"/api/public/cameras/{camera.id}/live-detection", headers=admin)
    assert resp.status_code == 200
    assert resp.json()["faces"] == [] and resp.json()["source"] == "kunlik"


async def test_start_run_conflict_and_cancel(client, admin, db_session):
    day = business_today() - timedelta(days=1)
    resp = await client.post("/api/video-tahlil/ishlar", json={"day": day.isoformat()}, headers=admin)
    assert resp.status_code == 201, resp.text
    run = resp.json()
    # Kamera yo'q — 0 vazifa, lekin tahlil navbatda.
    assert run["status"] == "navbatda" and run["jobsTotal"] == 0 and run["triggeredBy"] == "qolda"

    resp = await client.post("/api/video-tahlil/ishlar", json={"day": day.isoformat()}, headers=admin)
    assert resp.status_code == 409

    resp = await client.post(f"/api/video-tahlil/ishlar/{run['id']}/bekor", headers=admin)
    assert resp.status_code == 200 and resp.json()["status"] == "bekor"
    resp = await client.post(f"/api/video-tahlil/ishlar/{run['id']}/bekor", headers=admin)
    assert resp.status_code == 409

    actions = (await db_session.execute(select(AuditLog.action))).scalars().all()
    assert any("Video tahlil ishga tushirildi" in a for a in actions)
    assert any("bekor qilindi" in a for a in actions)


async def test_start_run_rejects_future_and_ancient_days(client, admin):
    future = business_today() + timedelta(days=1)
    assert (await client.post("/api/video-tahlil/ishlar", json={"day": future.isoformat()}, headers=admin)).status_code == 422
    old = business_today() - timedelta(days=90)
    assert (await client.post("/api/video-tahlil/ishlar", json={"day": old.isoformat()}, headers=admin)).status_code == 422


async def test_run_list_and_detail(client, admin, db_session):
    run = VideoAnalysisRun(day=date(2026, 10, 1), status="tugadi", window_start=datetime.now(timezone.utc),
                           window_end=datetime.now(timezone.utc), jobs_total=3, jobs_done=3)
    db_session.add(run)
    await db_session.commit()
    runs = (await client.get("/api/video-tahlil/ishlar", headers=admin)).json()
    assert runs[0]["id"] == str(run.id) and runs[0]["progress"] == 1.0
    detail = (await client.get(f"/api/video-tahlil/ishlar/{run.id}", headers=admin)).json()
    assert detail["jobs"] == [] and detail["errors"] == []
    assert (await client.get(f"/api/video-tahlil/ishlar/{uuid.uuid4()}", headers=admin)).status_code == 404


async def _results_fixture(db_session):
    day = date(2026, 10, 1)
    people = []
    for name, kind, group in [("Aliyev A", "talaba", "DI-1"), ("Boboyev B", "talaba", "DI-2"), ("Karimova K", "xodim", "Kafedra")]:
        person = StudentStaff(full_name=name, type=kind, group_or_position=group)
        db_session.add(person)
        people.append(person)
    await db_session.flush()
    db_session.add_all([
        DailyPersonCriteria(student_staff_id=people[0].id, day=day, attendance_status="kech_keldi", late_minutes=12,
                            coat_status="kiymagan", lessons_total=3, lessons_late=1, attention_score=55, smoking_events=1),
        DailyPersonCriteria(student_staff_id=people[1].id, day=day, attendance_status="kelmadi", coat_status="aniqlanmadi"),
        DailyPersonCriteria(student_staff_id=people[2].id, day=day, attendance_status="keldi", late_minutes=0,
                            early_leave="erta_ketdi", teacher_lessons=2, teacher_on_time=1, teacher_late=1,
                            teacher_activity=70, coat_status="kiygan"),
    ])
    await db_session.commit()
    return day, people


async def test_results_filters_search_and_paging(client, admin, db_session):
    day, people = await _results_fixture(db_session)
    url = "/api/video-tahlil/natijalar"
    body = (await client.get(url, params={"day": day.isoformat()}, headers=admin)).json()
    assert body["total"] == 3 and [r["fullName"] for r in body["items"]] == ["Aliyev A", "Boboyev B", "Karimova K"]
    # Kun berilmasa — oxirgi tahlil kuni.
    assert (await client.get(url, headers=admin)).json()["total"] == 3

    async def query(**params):
        resp = await client.get(url, params={"day": day.isoformat(), **params}, headers=admin)
        assert resp.status_code == 200, resp.text
        return [r["fullName"] for r in resp.json()["items"]]

    assert await query(type="xodim") == ["Karimova K"]
    assert await query(q="DI-2") == ["Boboyev B"]
    assert await query(filter="kech") == ["Aliyev A"]
    assert await query(filter="kelmadi") == ["Boboyev B"]
    assert await query(filter="xalatsiz") == ["Aliyev A"]
    assert await query(filter="chekish") == ["Aliyev A"]
    assert await query(filter="diqqat_past") == ["Aliyev A"]
    assert await query(filter="darsga_kech") == ["Aliyev A"]
    assert await query(filter="erta_ketdi") == ["Karimova K"]
    assert await query(filter="oqituvchi_kech") == ["Karimova K"]
    page = (await client.get(url, params={"day": day.isoformat(), "pageSize": 2, "page": 2}, headers=admin)).json()
    assert page["totalPages"] == 2 and [r["fullName"] for r in page["items"]] == ["Karimova K"]
    bad = await client.get(url, params={"filter": "yoq"}, headers=admin)
    assert bad.status_code == 422


async def test_results_summary_and_person_history(client, admin, db_session):
    day, people = await _results_fixture(db_session)
    summary = (await client.get("/api/video-tahlil/natijalar/xulosa", params={"day": day.isoformat()}, headers=admin)).json()
    assert summary["people"] == 3 and summary["present"] == 2 and summary["late"] == 1 and summary["absent"] == 1
    assert summary["coatNo"] == 1 and summary["coatYes"] == 1 and summary["coatUnknown"] == 1
    assert summary["smoking"] == 1 and summary["teacherLate"] == 1 and summary["teacherActivityAvg"] == 70
    assert summary["lessonsLate"] == 1 and summary["earlyLeave"] == 1 and summary["attentionAvg"] == 55
    history = (
        await client.get(f"/api/video-tahlil/natijalar/odam/{people[0].id}",
                         params={"from": "2026-09-01", "to": "2026-10-31"}, headers=admin)
    ).json()
    assert len(history) == 1 and history[0]["coatStatus"] == "kiymagan"
    assert (await client.get(f"/api/video-tahlil/natijalar/odam/{uuid.uuid4()}", headers=admin)).status_code == 404


async def test_operator_can_read_results_but_not_manage(client, operator, db_session):
    await _results_fixture(db_session)
    assert (await client.get("/api/video-tahlil/natijalar", headers=operator)).status_code == 200
    assert (await client.get("/api/video-tahlil/holat", headers=operator)).status_code == 200
    assert (await client.get("/api/video-tahlil/nvr", headers=operator)).status_code == 403
    day = (business_today() - timedelta(days=1)).isoformat()
    assert (await client.post("/api/video-tahlil/ishlar", json={"day": day}, headers=operator)).status_code == 403


async def test_anonymous_is_rejected(client, seeded):
    assert (await client.get("/api/video-tahlil/natijalar")).status_code == 401


async def test_nvr_crud_never_returns_password(client, admin, db_session):
    payload = {"name": "NVR-1", "ip": "192.168.0.94", "username": "admin", "password": "Sir12345", "maxStreams": 6}
    resp = await client.post("/api/video-tahlil/nvr", json=payload, headers=admin)
    assert resp.status_code == 201, resp.text
    nvr = resp.json()
    assert nvr["hasPassword"] is True and "password" not in nvr and nvr["fetchMode"] == "download"
    stored = await db_session.get(NvrDevice, uuid.UUID(nvr["id"]))
    assert stored.password and stored.password != "Sir12345"  # shifrlangan

    resp = await client.patch(f"/api/video-tahlil/nvr/{nvr['id']}", json={"stream": "sub", "password": ""}, headers=admin)
    assert resp.status_code == 200 and resp.json()["stream"] == "sub" and resp.json()["hasPassword"] is True

    bad = await client.post("/api/video-tahlil/nvr", json={"name": "x", "kind": "hikvision"}, headers=admin)
    assert bad.status_code == 422
    bad = await client.post("/api/video-tahlil/nvr", json={"name": "x", "kind": "fayl"}, headers=admin)
    assert bad.status_code == 422

    camera = Camera(name="K1", ip="192.168.0.11", zone="A", resolution="HD", status="faol")
    db_session.add(camera)
    await db_session.commit()
    resp = await client.put(f"/api/video-tahlil/kameralar/{camera.id}/nvr", json={"nvrId": nvr["id"], "channel": 4},
                            headers=admin)
    assert resp.status_code == 204
    assert (await client.put(f"/api/video-tahlil/kameralar/{camera.id}/nvr", json={"nvrId": nvr["id"]},
                             headers=admin)).status_code == 422
    listed = (await client.get("/api/video-tahlil/nvr", headers=admin)).json()
    assert listed[0]["cameras"] == 1

    assert (await client.delete(f"/api/video-tahlil/nvr/{nvr['id']}", headers=admin)).status_code == 204
    await db_session.refresh(camera)
    assert camera.nvr_id is None and camera.nvr_channel is None


async def test_folder_nvr_channels_and_auto_mapping(client, admin, db_session, tmp_path):
    for channel in ("1", "2", "7"):
        (tmp_path / channel).mkdir()
    (tmp_path / "not-a-channel").mkdir()
    nvr = (await client.post("/api/video-tahlil/nvr", json={"name": "Eksport", "kind": "fayl", "basePath": str(tmp_path)},
                             headers=admin)).json()
    test = (await client.post(f"/api/video-tahlil/nvr/{nvr['id']}/tekshirish", headers=admin)).json()
    assert test["ok"] is True and test["channels"] == 3
    channels = (await client.get(f"/api/video-tahlil/nvr/{nvr['id']}/kanallar", headers=admin)).json()
    assert [c["channel"] for c in channels] == [1, 2, 7]

    db_session.add(Camera(name="Papka 2", ip="10.9.9.9", zone="A", resolution="HD", status="faol"))
    await db_session.commit()
    mapped = (await client.post(f"/api/video-tahlil/nvr/{nvr['id']}/avto-boglash", headers=admin)).json()
    assert mapped["mapped"] == 1 and mapped["unmatchedChannels"] == [1, 7]
    channels = (await client.get(f"/api/video-tahlil/nvr/{nvr['id']}/kanallar", headers=admin)).json()
    assert next(c for c in channels if c["channel"] == 2)["cameraName"] == "Papka 2"


async def test_nvr_check_reports_failure(client, admin):
    nvr = (await client.post("/api/video-tahlil/nvr", json={"name": "Yo'q", "kind": "fayl", "basePath": "/yoq/papka"},
                             headers=admin)).json()
    test = (await client.post(f"/api/video-tahlil/nvr/{nvr['id']}/tekshirish", headers=admin)).json()
    assert test["ok"] is False and "Papka" in test["message"]
    listed = (await client.get("/api/video-tahlil/nvr", headers=admin)).json()
    assert listed[0]["lastError"]
    assert (await client.get(f"/api/video-tahlil/nvr/{nvr['id']}/kanallar", headers=admin)).status_code == 502


async def test_hikvision_auto_map_by_ip(client, admin, db_session, monkeypatch):
    from app.routers import video_analysis as router_module
    from app.services.nvr.isapi import NvrChannel

    class FakeIsapi:
        async def list_channels(self):
            return [NvrChannel(1, "Kirish", "192.168.0.11", True), NvrChannel(2, "201-xona", "192.168.0.99", True),
                    NvrChannel(3, "Noma'lum", None, False)]

    monkeypatch.setattr(router_module, "isapi_for", lambda nvr: FakeIsapi())
    db_session.add_all([
        Camera(name="Asosiy kirish", ip="192.168.0.11", zone="A", resolution="HD", status="faol"),
        Camera(name="201-xona", ip="192.168.0.12", zone="A", resolution="HD", status="faol"),
    ])
    await db_session.commit()
    nvr = (await client.post("/api/video-tahlil/nvr", json={"name": "NVR", "ip": "192.168.0.94"}, headers=admin)).json()
    mapped = (await client.post(f"/api/video-tahlil/nvr/{nvr['id']}/avto-boglash", headers=admin)).json()
    # 1 — IP bo'yicha, 2 — nom bo'yicha (IP o'zgargan), 3 — topilmadi.
    assert mapped["mapped"] == 2 and mapped["unmatchedChannels"] == [3]
    cams = {c.name: c.nvr_channel for c in (await db_session.execute(select(Camera))).scalars().unique()}
    assert cams == {"Asosiy kirish": 1, "201-xona": 2}


async def test_results_excel_export(client, admin, operator, db_session):
    from io import BytesIO

    from openpyxl import load_workbook

    day, people = await _results_fixture(db_session)
    resp = await client.get("/api/video-tahlil/natijalar/export.xlsx", params={"day": day.isoformat()}, headers=admin)
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "kunlik-tahlil-2026-10-01.xlsx" in resp.headers["content-disposition"]
    wb = load_workbook(BytesIO(resp.content))
    ws = wb["Natijalar"]
    names = [ws.cell(row=r, column=1).value for r in range(5, ws.max_row + 1)]
    assert names == ["Aliyev A", "Boboyev B", "Karimova K"]
    aliyev = [ws.cell(row=5, column=c).value for c in range(1, 21)]
    assert aliyev[3] == "Kech keldi" and aliyev[6] == 12 and aliyev[12] == "Kiymagan" and aliyev[14] == 1
    summary = {wb["Xulosa"].cell(row=r, column=1).value: wb["Xulosa"].cell(row=r, column=2).value for r in range(3, 18)}
    assert summary["Kelmadi"] == 1 and summary["Chekish holatlari"] == 1
    only_staff = await client.get("/api/video-tahlil/natijalar/export.xlsx",
                                  params={"day": day.isoformat(), "type": "xodim"}, headers=admin)
    ws = load_workbook(BytesIO(only_staff.content))["Natijalar"]
    assert ws.max_row == 5
    # Shaxsiy ma'lumot eksporti — faqat exportData huquqi bilan.
    assert (await client.get("/api/video-tahlil/natijalar/export.xlsx", headers=operator)).status_code == 403
