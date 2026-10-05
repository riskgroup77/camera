"""Hisobotlar (`/api/hisobot`) — tests/situation_world.py dunyosida.

Bugungi talaba davomati (faollar): keldi 6 (Botirova kechikib), kelmadi 2.
Davolash ishi: keldi 5, kelmadi 1. DI-2301 kecha+bugun: keldi 3, kelmadi 3.
Darslar (bugun): L1 — keldi 1, kech 1, kelmadi 2 (L2 hali baholanmagan).
O'qituvchilar (kecha+bugun): Yusupova L1 o'z vaqtida, LY kelmagan; Rahimov L2, L6 kelmagan.
"""

from datetime import datetime, time, timedelta, timezone

import pytest
from httpx import AsyncClient

from sqlalchemy import select

from app.models import DailyPersonCriteria, LessonAttendance, PresenceVisit
from app.batch import holatlar
from app.services import hisobot
from app.timezone import INSTITUTE_TZ
from tests.conftest import auth_headers
from tests.situation_world import _record, _situation_settings, world  # noqa: F401 — pytest fikstura


@pytest.fixture
async def admin(client: AsyncClient, world):
    return await auth_headers(client, "operator", "operator123")


async def _get(client, headers, **params):
    res = await client.get("/api/hisobot/report", params=params, headers=headers)
    assert res.status_code == 200, res.text
    return res.json()


def _ind(data, key):
    return next(c for c in data["criteria"] if c["key"] == key)["indicator"]


# ─────────────────────────────────────────── sof funksiyalar

def _m(name, raw, faculty=None):
    import uuid
    return hisobot.Member(uuid.uuid4(), name, raw, faculty, None, True)


def test_filter_members_student_cascade():
    import uuid
    fac = uuid.uuid4()
    ms = [_m("Ali", "2-kurs, DI-2301", fac), _m("Vali", "2-kurs, DI-2302", fac), _m("Soli", "3-kurs, DI-2101", fac),
          _m("Guli", "2-kurs, DI-2301", None)]
    f = hisobot.Filters(faculty=str(fac), course=2)
    assert [m.name for m in hisobot.filter_members(ms, "talaba", f)] == ["Ali", "Vali"]
    f = hisobot.Filters(faculty=str(fac), group=" di-2301 ")
    assert [m.name for m in hisobot.filter_members(ms, "talaba", f)] == ["Ali"]
    f = hisobot.Filters(faculty="none")
    assert [m.name for m in hisobot.filter_members(ms, "talaba", f)] == ["Guli"]
    f = hisobot.Filters(q="VAL")
    assert [m.name for m in hisobot.filter_members(ms, "talaba", f)] == ["Vali"]


def test_criteria_are_separate_per_kind():
    student = {c.key for c in hisobot.criteria_for("talaba")}
    staff = {c.key for c in hisobot.criteria_for("xodim")}
    # 2026-10-04: uxlash (#20) va ish vaqtidan tashqari kirish (#3) olib tashlandi;
    # oq xalat, chekish, diqqat, darsga kech kirish, darsdan erta chiqish va
    # o'qituvchi faolligi — kunlik video tahlildan.
    assert {"davomat", "kechikish", "dars_qatnashish", "darsga_kech", "darsdan_erta", "diqqat", "forma",
            "chekish"} == student
    assert {"davomat", "kechikish", "erta_ketish", "dars_otkazish", "faollik", "forma", "chekish"} == staff


def test_every_customer_criterion_has_a_report():
    """Buyurtmachi ro'yxatidagi 9 kriteriyaning har biri hisobotda bor va
    ro'yxat tartibida (6, 7, 8, 9, 10, 15, 19, 21, 22)."""
    codes = {hisobot.criterion_code(c.key, kind) for kind in hisobot.KINDS for c in hisobot.criteria_for(kind)}
    assert codes == {6, 7, 8, 9, 10, 15, 19, 21, 22}
    order = [hisobot.criterion_code(c.key, "talaba") for c in hisobot.criteria_for("talaba")]
    assert order == sorted(order)
    order = [hisobot.criterion_code(c.key, "xodim") for c in hisobot.criteria_for("xodim")]
    assert order == sorted(order)
    # Dalil turi bo'lgan mezonlar — holatlar ro'yxatidagi turlarga mos.
    from app.batch import holatlar

    for c in hisobot.CRITERIA:
        assert all(kind in holatlar.TYPES for kind in c.finding_types)


def test_teacher_rate_counts_only_on_time():
    t = hisobot.Tri(ok=1, late=1, miss=2)
    assert t.rate == 50.0  # talaba: kech kirgan ham kirgan
    assert t.on_time_rate == 25.0  # o'qituvchi: faqat o'z vaqtida
    assert t.rate_for("dars_otkazish") == 25.0 and t.rate_for("dars_qatnashish") == 50.0


# ─────────────────────────────────────────── API

async def test_student_attendance_all_and_faculty(client, admin, world):
    data = await _get(client, admin, kind="talaba")
    assert data["criterion"] == "davomat"
    assert data["population"]["total"] == 12  # nofaol hisobga kirmaydi
    assert _ind(data, "davomat") == "75%"
    assert _ind(data, "kechikish") == "1"
    assert data["report"]["breakdown"]["title"] == "Fakultetlar bo'yicha"
    # xodimlar mezoni talabada yo'q
    assert all(c["key"] not in ("erta_ketish", "dars_otkazish") for c in data["criteria"])

    data = await _get(client, admin, kind="talaba", faculty=str(world.di.id))
    rows = {r["name"]: r["value"] for r in data["report"]["breakdown"]["rows"]}
    assert data["report"]["breakdown"]["title"] == "Kurslar bo'yicha"
    assert rows["2-kurs"] == 75.0 and rows["3-kurs"] == 100.0


async def test_student_group_people_worst_first(client, admin, world):
    data = await _get(client, admin, kind="talaba", faculty=str(world.di.id), course=2, group="DI-2301",
                      **{"from": world.yesterday.isoformat(), "to": world.today.isoformat()})
    body = data["report"]
    assert _ind(data, "davomat") == "50%"
    names = [p["full_name"] for p in body["people"]]
    assert names[:3] == ["Choriyev Sardor", "Botirova Nigora", "Aliyev Anvar"]
    assert body["people"][1]["values"]["rate"] == 50.0
    assert len(body["trend"]["points"]) == 2


async def test_student_lessons_attention_and_early_exit(client, admin, world, db_session):
    """Diqqat va darsdan erta chiqish — kunlik video tahlil yozgan
    lesson_attendance ustunlaridan (app/batch/aggregate.py)."""
    rows = (await db_session.execute(select(LessonAttendance))).scalars().all()
    anvar = next(r for r in rows if r.student_staff_id == world.people.aliyev.id)
    anvar.attention_score, anvar.attention_samples, anvar.left_early = 35, 6, True
    await db_session.commit()

    data = await _get(client, admin, kind="talaba", criterion="dars_qatnashish")
    assert _ind(data, "dars_qatnashish") == "50%"
    assert _ind(data, "diqqat") == "1" and _ind(data, "darsdan_erta") == "1"
    data = await _get(client, admin, kind="talaba", criterion="diqqat")
    assert [p["full_name"] for p in data["report"]["people"]] == ["Aliyev Anvar"]
    assert data["report"]["people"][0]["values"]["events"] == 1
    data = await _get(client, admin, kind="talaba", criterion="darsdan_erta")
    assert [p["full_name"] for p in data["report"]["people"]] == ["Aliyev Anvar"]
    # Talaba diqqati xodimlarda ko'rinmaydi.
    staff = await _get(client, admin, kind="xodim")
    assert all(c["key"] not in ("diqqat", "darsdan_erta") for c in staff["criteria"])


async def test_coat_and_smoking_are_named_per_person(client, admin, world, db_session):
    """Oq xalat va chekish endi odamga bog'langan (kunlik natija qatori)."""
    db_session.add_all([
        DailyPersonCriteria(student_staff_id=world.people.aliyev.id, day=world.today, coat_status="kiymagan",
                            smoking_events=2),
        DailyPersonCriteria(student_staff_id=world.people.aliyev.id, day=world.yesterday, coat_status="kiymagan"),
        DailyPersonCriteria(student_staff_id=world.people.rahimov.id, day=world.today, coat_status="kiygan",
                            smoking_events=1),
    ])
    await db_session.commit()
    params = {"from": world.yesterday.isoformat(), "to": world.today.isoformat()}
    data = await _get(client, admin, kind="talaba", criterion="forma", **params)
    assert _ind(data, "forma") == "2" and _ind(data, "chekish") == "2"
    assert [p["full_name"] for p in data["report"]["people"]] == ["Aliyev Anvar"]
    assert data["report"]["people"][0]["values"]["events"] == 2
    assert {p["value"] for p in data["report"]["trend"]["points"]} == {1}
    staff = await _get(client, admin, kind="xodim", criterion="chekish", **params)
    assert _ind(staff, "forma") == "0" and _ind(staff, "chekish") == "1"
    assert [p["full_name"] for p in staff["report"]["people"]] == ["Rahimov Bobur"]
    assert any("chekish" in line for line in staff["report"]["summary"])


async def test_staff_punctuality_lateness_and_units(client, admin, world):
    params = {"from": world.yesterday.isoformat(), "to": world.today.isoformat()}
    data = await _get(client, admin, kind="xodim", criterion="dars_otkazish", **params)
    assert _ind(data, "dars_otkazish") == "25%"
    assert [p["full_name"] for p in data["report"]["people"]] == ["Rahimov Bobur", "Yusupova Dilnoza Anvarovna"]

    data = await _get(client, admin, kind="xodim", criterion="kechikish")
    assert data["report"]["people"][0]["full_name"] == "Karimov Aziz Olimovich"
    assert data["report"]["people"][0]["values"]["late_min"] == 80  # bitta kun — daqiqada

    options = (await client.get("/api/hisobot/filters", params={"kind": "xodim"}, headers=admin)).json()
    anatomy = next(u for u in options["units"] if u["name"] == "Anatomiya kafedrasi")
    assert anatomy["count"] == 2
    data = await _get(client, admin, kind="xodim", unit=anatomy["id"])
    assert data["population"]["total"] == 2
    data = await _get(client, admin, kind="xodim", unit_kind="lavozim")
    assert data["population"]["total"] == 1


async def test_lesson_lateness_activity_and_evidence_per_person(client, admin, world, db_session):
    """#8 darsga kech kirish, #21 o'qituvchi faolligi, #22 kech/kirmadi alohida
    va har odam qatorida shu mezon bo'yicha video dalillar soni."""
    from app.models import LessonSession

    lessons = (await db_session.execute(select(LessonSession))).scalars().all()
    rahimov_lesson = next(row for row in lessons if row.teacher_id == world.people.rahimov.id
                          and row.teacher_on_time is False)
    # Rahimov bitta darsga kech KIRDI (xonada ko'rindi), faolligi past.
    rahimov_lesson.teacher_first_seen_at = datetime.combine(rahimov_lesson.date, time(9, 25), tzinfo=INSTITUTE_TZ)
    rahimov_lesson.teacher_activity_score, rahimov_lesson.activity_samples = 20, 4
    # O'lchanmagan (activity_samples=0) dars 0 ball bilan bo'lsa ham "past" emas.
    other = next(row for row in lessons if row.id != rahimov_lesson.id)
    other.teacher_activity_score, other.activity_samples = 0, 0
    clip = {"kod": 22, "tur": "oqituvchi_kech", "sabab": "Darsga kech kirdi", "kamera_id": None,
            "vaqt": rahimov_lesson.teacher_first_seen_at.isoformat(), "klip": "video-tahlil/dalil/1.mp4"}
    no_clip = {"kod": 21, "tur": "faollik_past", "sabab": "Faollik past", "kamera_id": None,
               "vaqt": rahimov_lesson.teacher_first_seen_at.isoformat(), "klip": None, "klip_xato": "yozuv yo'q"}
    db_session.add(DailyPersonCriteria(student_staff_id=world.people.rahimov.id, day=rahimov_lesson.date,
                                       details={"dalillar": [clip, no_clip]}))
    anvar = next(r for r in (await db_session.execute(select(LessonAttendance))).scalars().all()
                 if r.student_staff_id == world.people.aliyev.id)
    anvar.status = "kech_keldi"
    await db_session.commit()
    params = {"from": world.yesterday.isoformat(), "to": world.today.isoformat()}

    staff = await _get(client, admin, kind="xodim", criterion="dars_otkazish", **params)
    by_key = {c["key"]: c for c in staff["criteria"]}
    assert by_key["dars_otkazish"]["code"] == 22 and by_key["faollik"]["code"] == 21
    assert by_key["kechikish"]["code"] == 6
    rahimov = next(p for p in staff["report"]["people"] if p["full_name"] == "Rahimov Bobur")
    assert rahimov["values"]["late"] == 1 and rahimov["values"]["miss"] == 1
    assert rahimov["values"]["dalil"] == 1  # faqat video bor holatlar
    assert [c["key"] for c in staff["report"]["columns"]][-1] == "dalil"
    labels = {t["label"]: t["value"] for t in staff["report"]["tiles"]}
    assert labels["O'qituvchi kech kirgan"] == 1

    staff = await _get(client, admin, kind="xodim", criterion="faollik", **params)
    assert _ind(staff, "faollik") == "1"
    assert [p["full_name"] for p in staff["report"]["people"]] == ["Rahimov Bobur"]
    assert staff["report"]["people"][0]["values"]["dalil"] is None  # klip kesilmagan

    students = await _get(client, admin, kind="talaba", criterion="darsga_kech", **params)
    by_key = {c["key"]: c for c in students["criteria"]}
    assert by_key["darsga_kech"]["code"] == 8 and by_key["kechikish"]["code"] == 7
    assert "Aliyev Anvar" in [p["full_name"] for p in students["report"]["people"]]


async def test_staff_early_leave_and_coat(client, admin, world, db_session):
    today = world.today
    monday = today - timedelta(days=today.weekday() + 7)  # o'tgan haftaning dushanbasi
    db_session.add(_record(world.people.rahimov, monday, "keldi", "07:50", "15:00"))
    # "Erta ketdi" deyish uchun dalil kerak: o'sha kuni odam kameralarda bir
    # necha marta ko'rilgan va oxirgi ko'rinish aynan 15:00 (app/services/
    # attendance_policy.early_leave_verdict).
    db_session.add(
        PresenceVisit(
            student_staff_id=world.people.rahimov.id,
            first_seen_at=datetime.combine(monday, time(7, 50), tzinfo=INSTITUTE_TZ),
            last_seen_at=datetime.combine(monday, time(15, 0), tzinfo=INSTITUTE_TZ),
            sightings=9,
        )
    )
    db_session.add(DailyPersonCriteria(student_staff_id=world.people.rahimov.id, day=monday, coat_status="kiymagan"))
    await db_session.commit()
    params = {"from": monday.isoformat(), "to": today.isoformat()}
    data = await _get(client, admin, kind="xodim", criterion="erta_ketish", **params)
    assert _ind(data, "erta_ketish") == "1"
    assert data["report"]["people"][0]["full_name"] == "Rahimov Bobur"

    data = await _get(client, admin, kind="xodim", criterion="forma", **params)
    body = data["report"]
    assert [p["full_name"] for p in body["people"]] == ["Rahimov Bobur"]
    assert body["columns"][0]["key"] == "events"
    assert sum(r["value"] or 0 for r in body["breakdown"]["rows"]) == 1


async def test_student_filters_and_export(client, admin, world):
    options = (await client.get("/api/hisobot/filters", params={"kind": "talaba"}, headers=admin)).json()
    assert {f["name"] for f in options["faculties"]} >= {"Davolash ishi", "Pediatriya", "Fakultetsiz"}
    assert any(g["name"] == "DI-2301" and g["course"] == 2 for g in options["groups"])

    res = await client.get("/api/hisobot/export.xlsx", params={"kind": "talaba", "criterion": "davomat"},
                           headers=admin)
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert res.content[:2] == b"PK"


# ─────────────────────────────────────────── tushunarlilik: gap, izoh, bo'sh holat

async def test_summary_sentence_matches_the_tiles(client, admin, world):
    """Tepadagi gap plitkalardagi ayni sonlardan tuziladi."""
    data = await _get(client, admin, kind="talaba")
    body = data["report"]
    assert data["scope"] == "Barcha talabalar"
    first = body["summary"][0]
    # bugun: 6 keldi (1 tasi kech), 2 kelmadi, 4 kishida yozuv yo'q
    assert "12 talabadan" in first and "bugun" in first
    assert "5 tasi o'z vaqtida keldi" in first
    assert "1 tasi kech keldi" in first
    assert "2 tasi kelmadi" in first
    assert "4 tasi haqida yozuv yo'q" in first
    assert "Kelganlar ulushi — 75%." in body["summary"]
    # yuzi kiritilmaganlar haqida ogohlantirish (12 dan 3 tasi)
    assert any("3 tasining yuzi tizimga kiritilmagan" in line for line in body["summary"])
    assert data["population"] == {"total": 12, "enrolled": 9, "not_enrolled": 3}


async def test_criteria_explain_themselves(client, admin, world):
    data = await _get(client, admin, kind="xodim")
    by_key = {c["key"]: c for c in data["criteria"]}
    assert all(c["description"] for c in data["criteria"])
    # vaqtlar kodda emas, ish vaqti sozlamasidan (08:00 + 10 daqiqa)
    assert "08:10" in by_key["kechikish"]["description"]
    assert "17:00" in by_key["erta_ketish"]["description"]
    assert all(c["available"] is (c["unavailable"] is None) for c in data["criteria"])


async def test_blocked_criterion_says_why_instead_of_zero(client, admin, world):
    """Dars jadvali yo'q kunda — "0%" emas, sabab."""
    old = (world.today - timedelta(days=60)).isoformat()
    data = await _get(client, admin, kind="talaba", criterion="dars_qatnashish", **{"from": old, "to": old})
    assert _ind(data, "dars_qatnashish") == "—"
    body = data["report"]
    assert body["blocked"] is True
    assert "dars jadvali" in body["note"]
    assert body["empty"]["description"]
    assert body["summary"][1] == body["note"]


async def test_day_table_has_reader_friendly_columns(client, admin, world):
    """Bitta kun tanlanganda jadval: holat, vaqtlar, kechikish, izoh."""
    data = await _get(client, admin, kind="talaba")
    body = data["report"]
    assert [c["key"] for c in body["columns"]] == ["holat", "check_in", "check_out", "late_min", "note"]
    assert [c["label"] for c in body["columns"]][:2] == ["Holati", "Kelgan vaqti"]
    assert body["people_total"] == 12  # yozuvi yo'qlar ham ko'rinadi
    rows = {p["full_name"]: p["values"] for p in body["people"]}
    assert rows["Botirova Nigora"]["holat"] == "Kech keldi"
    assert rows["Botirova Nigora"]["check_in"] == "09:15"
    assert rows["Botirova Nigora"]["late_min"] == 75
    assert rows["Choriyev Sardor"]["holat"] == "Kelmadi"
    # yuzi kiritilmagan odam uchun sabab yoziladi, bo'sh katak emas
    assert "Yuzi tizimga kiritilmagan" in rows["Ergasheva Laylo"]["note"]
    assert body["people"][0]["values"]["holat"] == "Kelmadi"  # eng muammolisi birinchi
    assert "Tartib:" in body["people_hint"]


async def test_export_columns_match_the_table(client, admin, world):
    from io import BytesIO

    from openpyxl import load_workbook

    data = await _get(client, admin, kind="talaba")
    res = await client.get("/api/hisobot/export.xlsx", params={"kind": "talaba"}, headers=admin)
    assert res.status_code == 200
    wb = load_workbook(BytesIO(res.content))
    rows = list(wb["Odamlar"].iter_rows(values_only=True))
    header = [c for c in rows[3] if c is not None]
    expected = ["№", "F.I.Sh.", "Guruh yoki bo'linma"] + [
        f"{c['label']}, {c['unit']}" if c["unit"] else c["label"] for c in data["report"]["columns"]]
    assert header == expected
    assert wb["Hisobot"]["B2"].value == data["scope"]


async def test_requires_auth(client, world):
    res = await client.get("/api/hisobot/report", params={"kind": "talaba"})
    assert res.status_code == 401




# ─────────────────────────────────────────── guruh: talaba × mezon (Nazorat)

def test_matrix_cell_reads_like_a_register():
    import uuid
    from datetime import date

    pid = uuid.uuid4()
    day = date(2026, 10, 6)
    data = hisobot.Data("talaba", day, day, None, [])
    data.day_rows[pid] = {"status": "kech_keldi", "check_in": time(8, 42), "late": 32}
    data.lessons[pid] = hisobot.Tri(ok=1, late=1, miss=1)
    data.events["forma"] = {pid: 1}
    data.findings[holatlar.NO_COAT] = {pid: (1, 1)}
    cell = hisobot.matrix_cell
    assert cell(data, "davomat", pid)["value"] == "kech" and cell(data, "davomat", pid)["tone"] == "warning"
    assert cell(data, "kechikish", pid)["value"] == "+32 daq"
    assert cell(data, "dars_qatnashish", pid)["value"] == "2/3"
    assert cell(data, "forma", pid)["value"] == "1" and cell(data, "forma", pid)["tone"] == "danger"
    assert cell(data, "forma", pid)["evidence"] == 1  # 2 daqiqalik video dalil
    assert cell(data, "chekish", pid)["value"] == "0" and cell(data, "chekish", pid)["evidence"] is None
    other = uuid.uuid4()
    assert cell(data, "davomat", other)["value"] == "—"
    assert cell(data, "dars_qatnashish", other)["value"] == "—"
    # Bir necha kun: kelgan kunlar / yozuvli kunlar.
    week = hisobot.Data("talaba", day - timedelta(days=4), day, None, [])
    week.att[pid] = hisobot.Att(present=3, late=1, absent=2)
    assert cell(week, "davomat", pid)["value"] == "3/5" and cell(week, "kechikish", pid)["value"] == "1"


async def test_group_matrix_lists_every_student_with_every_criterion(client, admin, world, db_session):
    res = await client.get("/api/situation/group-criteria", params={"group": "DI-2301"}, headers=admin)
    assert res.status_code == 200, res.text
    body = res.json()
    names = [p["full_name"] for p in body["people"]]
    # Faol a'zolar alifbo tartibida; XDI-2301 (Komilov) va faol emaslar yo'q.
    assert names == ["Aliyev Anvar", "Botirova Nigora", "Choriyev Sardor", "Davronov Jasur", "Ergasheva Laylo"]
    keys = [c["key"] for c in body["criteria"]]
    assert keys[:2] == ["davomat", "kechikish"] and {"forma", "chekish", "diqqat"} <= set(keys)
    cells = {p["full_name"]: p["cells"] for p in body["people"]}
    assert cells["Aliyev Anvar"]["davomat"]["value"] == "keldi"
    assert cells["Botirova Nigora"]["davomat"]["value"] == "kech"
    assert cells["Botirova Nigora"]["kechikish"]["value"].startswith("+")
    assert cells["Choriyev Sardor"]["davomat"]["value"] == "kelmadi"
    # Video tahlil o'tmagan kun: oq xalat "0" emas — sababi bilan "—".
    forma = next(c for c in body["criteria"] if c["key"] == "forma")
    assert forma["unavailable"] and "tahlil" in forma["unavailable"]
    assert cells["Aliyev Anvar"]["forma"]["value"] == "—"

    db_session.add(DailyPersonCriteria(student_staff_id=world.people.aliyev.id, day=world.today, coat_status="kiymagan"))
    await db_session.commit()
    body = (await client.get("/api/situation/group-criteria", params={"group": "DI-2301"}, headers=admin)).json()
    cells = {p["full_name"]: p["cells"] for p in body["people"]}
    assert body["analysed"] is True
    assert cells["Aliyev Anvar"]["forma"]["value"] == "1" and cells["Botirova Nigora"]["forma"]["value"] == "0"


async def test_group_matrix_requires_login_but_not_the_report_password(client, admin, world, monkeypatch):
    assert (await client.get("/api/situation/group-criteria", params={"group": "DI-2301"})).status_code in (401, 403)
    # Hisobotlar paroli yoqilgan bo'lsa ham Nazorat kriteriyalari ochiladi.
    from app.config import settings
    monkeypatch.setattr(settings, "report_password_hash", "qulf")
    res = await client.get("/api/situation/group-criteria", params={"group": "MD-134/25"}, headers=admin)
    assert res.status_code == 200 and res.json()["people"] == []
