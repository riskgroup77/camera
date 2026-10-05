"""Hisobot: kriteriya kartalari -> ro'yxat -> odam kesimi.

Hisobotning eng katta xavfi — bo'sh raqamni "hammasi joyida" deb
ko'rsatish. Shuning uchun bu testlar raqamlar to'g'riligidan tashqari
IZOHLAR ham chiqishini tekshiradi: modul o'chirilgan yoki yuzlar
ro'yxatga kiritilmagan bo'lsa, karta buni aytishi shart.
"""

from datetime import date, datetime, time, timedelta, timezone

import pytest
from httpx import AsyncClient
from sqlalchemy import select

from app.models import AIModuleConfig, AttendanceRecord, Building, Camera, Event, Faculty, PresenceVisit, StudentStaff
from app.timezone import local_now
from tests.conftest import auth_headers


def _utc(day: date, hour: int) -> datetime:
    """Toshkent vaqtidagi soatni UTC'ga o'tkazadi (tashriflar UTC'da)."""
    from app.timezone import INSTITUTE_TZ

    return datetime.combine(day, time(hour, 0), tzinfo=INSTITUTE_TZ).astimezone(timezone.utc)


@pytest.fixture
async def people(db_session, seeded):
    faculty = (await db_session.execute(select(Faculty))).scalars().first()
    rows = {
        "came": StudentStaff(
            full_name="Kelgan Xodim", type="xodim", faculty_id=faculty.id,
            group_or_position="Assistent", biometrics_status="tasdiqlangan",
        ),
        "late": StudentStaff(
            full_name="Kechikkan Xodim", type="xodim", faculty_id=faculty.id,
            group_or_position="Dotsent", biometrics_status="tasdiqlangan",
        ),
        "absent": StudentStaff(
            full_name="Kelmagan Xodim", type="xodim", faculty_id=faculty.id,
            group_or_position="Professor", biometrics_status="tasdiqlangan",
        ),
        "unenrolled": StudentStaff(
            full_name="Royxatsiz Xodim", type="xodim", faculty_id=faculty.id,
            group_or_position="Laborant", biometrics_status="yoq",
        ),
        "student": StudentStaff(
            full_name="Sinov Talaba", type="talaba", faculty_id=faculty.id,
            group_or_position="2-kurs, DI-1625", biometrics_status="yoq",
        ),
    }
    db_session.add_all(list(rows.values()))
    await db_session.commit()
    for row in rows.values():
        await db_session.refresh(row)
    return rows


@pytest.fixture
async def today_data(db_session, people):
    today = local_now().date()
    camera = (await db_session.execute(select(Camera))).scalars().first()
    if camera is None:
        building = (await db_session.execute(select(Building))).scalars().first()
        camera = Camera(
            name="Kirish-H", ip="10.5.0.1", building_id=building.id, zone="Kirish",
            resolution="1080p", status="faol",
        )
        db_session.add(camera)
        await db_session.commit()
        await db_session.refresh(camera)

    db_session.add_all(
        [
            AttendanceRecord(
                student_staff_id=people["came"].id, date=today, status="keldi", check_in=time(8, 40)
            ),
            AttendanceRecord(
                student_staff_id=people["late"].id, date=today, status="kech_keldi", check_in=time(9, 35)
            ),
            AttendanceRecord(student_staff_id=people["absent"].id, date=today, status="kelmadi"),
            PresenceVisit(
                student_staff_id=people["came"].id,
                camera_id=camera.id,
                first_seen_at=_utc(today, 8),
                last_seen_at=_utc(today, 8) + timedelta(minutes=2),
                sightings=3,
            ),
        ]
    )
    await db_session.commit()
    return camera


@pytest.mark.usefixtures("seeded")
class TestCriteriaCards:
    async def test_attendance_card_counts_each_status(self, client: AsyncClient, today_data):
        headers = await auth_headers(client, "admin", "admin123")
        body = (
            await client.get("/api/reports/criteria", headers=headers, params={"population": "xodim", "period": "bugun"})
        ).json()

        assert body["period"]["key"] == "bugun" and body["period"]["days"] == 1
        davomat = next(item for item in body["criteria"] if item["key"] == "davomat")
        counts = {bucket["key"]: bucket["count"] for bucket in davomat["buckets"]}
        assert counts == {"keldi": 1, "kech_keldi": 1, "kelmadi": 1}
        assert davomat["note"] is None

    async def test_seen_card_separates_recognised_from_unseen(self, client: AsyncClient, today_data):
        headers = await auth_headers(client, "admin", "admin123")
        body = (
            await client.get("/api/reports/criteria", headers=headers, params={"population": "xodim", "period": "bugun"})
        ).json()
        korinish = next(item for item in body["criteria"] if item["key"] == "korinish")
        counts = {bucket["key"]: bucket["count"] for bucket in korinish["buckets"]}
        # Uch kishi ro'yxatdan o'tgan, kamerada bittasi ko'rindi.
        assert counts["korindi"] == 1 and counts["korinmadi"] == 2

    async def test_enrolment_card_shows_current_coverage(self, client: AsyncClient, people):
        headers = await auth_headers(client, "admin", "admin123")
        body = (
            await client.get("/api/reports/criteria", headers=headers, params={"population": "xodim", "period": "bugun"})
        ).json()
        biometrika = next(item for item in body["criteria"] if item["key"] == "biometrika")
        counts = {bucket["key"]: bucket["count"] for bucket in biometrika["buckets"]}
        assert counts["tasdiqlangan"] == 3 and counts["tasdiqlanmagan"] == 1

    async def test_students_card_explains_why_it_is_empty(self, client: AsyncClient, db_session, people):
        """#7 o'chirilgan: nol raqam "hech kim kelmadi" degani emas."""
        module = (
            await db_session.execute(select(AIModuleConfig).where(AIModuleConfig.code == 7))
        ).scalar_one()
        module.active = False  # 2026-10-04 dan standart bo'yicha yoqiq
        await db_session.commit()
        headers = await auth_headers(client, "admin", "admin123")
        body = (
            await client.get(
                "/api/reports/criteria", headers=headers, params={"population": "talaba", "period": "bugun"}
            )
        ).json()
        davomat = next(item for item in body["criteria"] if item["key"] == "davomat")
        assert davomat["note"] and "o'chirilgan" in davomat["note"]

        module = (
            await db_session.execute(select(AIModuleConfig).where(AIModuleConfig.code == 7))
        ).scalar_one()
        module.active = True
        await db_session.commit()
        again = (
            await client.get(
                "/api/reports/criteria", headers=headers, params={"population": "talaba", "period": "bugun"}
            )
        ).json()
        davomat_again = next(item for item in again["criteria"] if item["key"] == "davomat")
        # Endi sabab boshqacha: modul yoqilgan, lekin yuzlar ro'yxatda yo'q.
        assert davomat_again["note"] and "yuzi" in davomat_again["note"]

    async def test_event_criteria_carry_module_codes(self, client: AsyncClient, db_session, people):
        # Chekish (#15) sinov rejimida ham kartada sanaladi — operator
        # uni «Tekshiruv»da tasdiqlaydi.
        db_session.add(
            Event(
                camera_name="Hovli", building="1-bino", module_code=15,
                module_name="Chekish / elektron sigareta", group="D", confidence=70,
                severity="o'rta", status="yangi", is_trial=True,
            )
        )
        await db_session.commit()
        headers = await auth_headers(client, "admin", "admin123")
        body = (
            await client.get("/api/reports/criteria", headers=headers, params={"population": "xodim", "period": "bugun"})
        ).json()
        keys = {item["key"] for item in body["criteria"]}
        assert "off_hours" not in keys and "uxlash" not in keys
        smoking = next(item for item in body["criteria"] if item["key"] == "chekish")
        assert smoking["detail"] == "events" and smoking["moduleCodes"] == [15]
        assert smoking["total"] == 1
        punctuality = next(item for item in body["criteria"] if item["key"] == "dars_intizomi")
        assert punctuality["moduleCodes"] == [22]

    async def test_yesterday_is_a_different_window(self, client: AsyncClient, today_data):
        headers = await auth_headers(client, "admin", "admin123")
        body = (
            await client.get("/api/reports/criteria", headers=headers, params={"population": "xodim", "period": "kecha"})
        ).json()
        davomat = next(item for item in body["criteria"] if item["key"] == "davomat")
        assert davomat["total"] == 0
        assert body["period"]["start"] == body["period"]["end"]


@pytest.mark.usefixtures("seeded")
class TestCriterionPeople:
    async def test_list_behind_a_bucket(self, client: AsyncClient, today_data, people):
        headers = await auth_headers(client, "admin", "admin123")
        body = (
            await client.get(
                "/api/reports/criteria/davomat/people",
                headers=headers,
                params={"population": "xodim", "period": "bugun", "bucket": "kelmadi"},
            )
        ).json()
        assert [item["fullName"] for item in body["items"]] == ["Kelmagan Xodim"]
        assert body["items"][0]["absentDays"] == 1

    async def test_row_carries_the_proof_of_arrival(self, client: AsyncClient, today_data, people):
        headers = await auth_headers(client, "admin", "admin123")
        body = (
            await client.get(
                "/api/reports/criteria/davomat/people",
                headers=headers,
                params={"population": "xodim", "period": "bugun", "bucket": "keldi"},
            )
        ).json()
        row = next(item for item in body["items"] if item["fullName"] == "Kelgan Xodim")
        assert row["firstCheckIn"] == "08:40"
        assert row["visits"] == 1 and row["cameras"] == 1
        assert row["lastSeenCamera"]  # qaysi kamera ko'rgani — isbot

    async def test_unseen_list_only_has_enrolled_people(self, client: AsyncClient, today_data, people):
        headers = await auth_headers(client, "admin", "admin123")
        body = (
            await client.get(
                "/api/reports/criteria/korinish/people",
                headers=headers,
                params={"population": "xodim", "period": "bugun", "bucket": "korinmadi"},
            )
        ).json()
        names = [item["fullName"] for item in body["items"]]
        assert "Kelmagan Xodim" in names
        assert "Royxatsiz Xodim" not in names  # yuzi yo'q — kamera taniy olmaydi

    async def test_search_filters_the_list(self, client: AsyncClient, today_data, people):
        headers = await auth_headers(client, "admin", "admin123")
        body = (
            await client.get(
                "/api/reports/criteria/biometrika/people",
                headers=headers,
                params={"population": "xodim", "period": "bugun", "bucket": "tasdiqlangan", "search": "kechikkan"},
            )
        ).json()
        assert [item["fullName"] for item in body["items"]] == ["Kechikkan Xodim"]


@pytest.mark.usefixtures("seeded")
class TestPersonDetail:
    async def test_detail_has_identity_period_stats_and_days(self, client: AsyncClient, today_data, people):
        headers = await auth_headers(client, "admin", "admin123")
        person = people["came"]
        body = (
            await client.get(
                f"/api/reports/people/{person.id}", headers=headers, params={"period": "hafta"}
            )
        ).json()

        assert body["fullName"] == "Kelgan Xodim"
        assert body["unit"] == "Assistent" and body["faculty"]
        assert body["period"]["days"] == 7 and len(body["days"]) == 7
        assert body["presentDays"] == 1 and body["visits"] == 1
        today_row = next(row for row in body["days"] if row["date"] == local_now().date().isoformat())
        assert today_row["status"] == "keldi" and today_row["checkIn"] == "08:40"
        assert today_row["visits"] == 1 and today_row["firstCamera"]

    async def test_unenrolled_person_is_explained_not_blamed(self, client: AsyncClient, people):
        headers = await auth_headers(client, "admin", "admin123")
        body = (
            await client.get(
                f"/api/reports/people/{people['unenrolled'].id}", headers=headers, params={"period": "bugun"}
            )
        ).json()
        assert body["note"] and "ro'yxatga kiritilmagan" in body["note"]
        assert body["presentDays"] == 0

    async def test_unknown_person_is_404(self, client: AsyncClient):
        headers = await auth_headers(client, "admin", "admin123")
        assert (
            await client.get(
                "/api/reports/people/00000000-0000-0000-0000-000000000000", headers=headers
            )
        ).status_code == 404
        assert (await client.get("/api/reports/people/not-a-uuid", headers=headers)).status_code == 404


def _workbook(content: bytes):
    from io import BytesIO

    from openpyxl import load_workbook

    return load_workbook(BytesIO(content))


@pytest.mark.usefixtures("seeded")
class TestCriteriaExport:
    """Hisobotlar sahifasidan Excel: qisqa — raqamlar, to'liq — ism-familiyalar bilan."""

    async def test_short_file_has_only_the_numbers(self, client: AsyncClient, today_data):
        headers = await auth_headers(client, "admin", "admin123")
        response = await client.get(
            "/api/reports/criteria.xlsx",
            headers=headers,
            params={"population": "xodim", "period": "bugun", "variant": "qisqa"},
        )
        assert response.status_code == 200
        assert "qisqa" in response.headers["content-disposition"]
        wb = _workbook(response.content)
        assert wb.sheetnames == ["Xulosa"]
        rows = [tuple(r) for r in wb["Xulosa"].iter_rows(values_only=True)]
        assert ("Davomat", "Keldi", 1) in [tuple(r[:3]) for r in rows]
        assert ("", "Kelmadi", 1) in [tuple(r[:3]) for r in rows] or (None, "Kelmadi", 1) in [tuple(r[:3]) for r in rows]
        names = " ".join(str(v) for r in rows for v in r if v)
        assert "Kelgan Xodim" not in names

    async def test_full_file_lists_people_by_name(self, client: AsyncClient, today_data):
        headers = await auth_headers(client, "admin", "admin123")
        response = await client.get(
            "/api/reports/criteria.xlsx",
            headers=headers,
            params={"population": "xodim", "period": "bugun", "variant": "toliq"},
        )
        assert response.status_code == 200
        wb = _workbook(response.content)
        assert "Davomat" in wb.sheetnames
        attendance = [r for r in wb["Davomat"].iter_rows(min_row=2, values_only=True)]
        names = {r[2]: r for r in attendance}
        assert names["Kelgan Xodim"][5] == "Keldi"
        assert names["Kelgan Xodim"][6] == "08:40"
        assert names["Kechikkan Xodim"][5] == "Kechikdi"

    async def test_full_file_follows_the_selected_bucket(self, client: AsyncClient, today_data):
        headers = await auth_headers(client, "admin", "admin123")
        response = await client.get(
            "/api/reports/criteria.xlsx",
            headers=headers,
            params={
                "population": "xodim", "period": "bugun", "variant": "toliq",
                "criterion": "davomat", "bucket": "keldi",
            },
        )
        wb = _workbook(response.content)
        assert wb.sheetnames == ["Xulosa", "Davomat"]
        people = [r[2] for r in wb["Davomat"].iter_rows(min_row=2, values_only=True)]
        assert people == ["Kelgan Xodim"]
